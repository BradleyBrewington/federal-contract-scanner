"""Export opportunities to a formatted Excel workbook."""

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

logger = logging.getLogger(__name__)

# --- Style constants ---
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

BODY_FONT = Font(name="Calibri", size=10)
WRAP_ALIGN = Alignment(vertical="top", wrap_text=True)
TOP_ALIGN = Alignment(vertical="top")
LINK_FONT = Font(name="Calibri", size=10, color="0563C1", underline="single")

THIN_BORDER = Border(
    bottom=Side(style="thin", color="D9D9D9"),
)

SCORE_GREEN = PatternFill("solid", fgColor="C6EFCE")
SCORE_YELLOW = PatternFill("solid", fgColor="FFEB9C")
SCORE_RED = PatternFill("solid", fgColor="FFC7CE")
SCORE_GREEN_FONT = Font(name="Calibri", size=12, bold=True, color="006100")
SCORE_YELLOW_FONT = Font(name="Calibri", size=12, bold=True, color="9C6500")
SCORE_RED_FONT = Font(name="Calibri", size=12, bold=True, color="9C0006")

STRIPE_FILL = PatternFill("solid", fgColor="F2F7FB")


def export_to_excel(opportunities, output_path):
    """Write opportunities to a two-sheet Excel workbook.

    Sheet 1 — Dashboard:  one-line-per-opp for fast scanning.
    Sheet 2 — Details:    full info including descriptions, contacts, AI notes.
    """
    wb = Workbook()

    _build_dashboard(wb, opportunities)
    _build_details(wb, opportunities)

    wb.save(output_path)
    logger.info(f"Excel report saved: {output_path} ({len(opportunities)} rows)")
    return output_path


# ------------------------------------------------------------------ dashboard
DASH_COLUMNS = [
    ("Score",       6),
    ("Title",       52),
    ("Agency",      28),
    ("Type",        22),
    ("Set-Aside",   22),
    ("NAICS",       8),
    ("Posted",      11),
    ("First Seen",  11),
    ("Deadline",    11),
    ("Days Left",   9),
    ("Flags",       40),
    ("SAM.gov Link", 14),
]


def _build_dashboard(wb, opportunities):
    ws = wb.active
    ws.title = "Dashboard"
    ws.sheet_properties.tabColor = "1F4E79"

    # Header row
    for col_idx, (header, width) in enumerate(DASH_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

    # Data rows
    now = datetime.now()
    for row_idx, opp in enumerate(opportunities, 2):
        score = opp.get("score", 0)
        deadline_str = opp.get("response_deadline") or ""
        days_left = _days_until(deadline_str, now)
        flags_text = " | ".join(opp.get("flags") or [])

        values = [
            score,
            opp.get("title", ""),
            opp.get("agency", ""),
            opp.get("type", ""),
            opp.get("set_aside", ""),
            opp.get("naics_code", ""),
            _format_date_short(opp.get("posted_date", "")),
            _format_date_short(opp.get("first_seen", "")),
            _format_date_short(deadline_str),
            days_left if days_left is not None else "",
            flags_text,
            opp.get("url", ""),
        ]

        for col_idx, val in enumerate(values, 1):
            # Strip non-printable characters that Excel rejects
            if isinstance(val, str):
                val = "".join(c for c in val if c >= " " or c in "\t\n\r")
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = BODY_FONT
            cell.alignment = TOP_ALIGN
            cell.border = THIN_BORDER

        # Alternate row shading
        if row_idx % 2 == 0:
            for col_idx in range(1, len(DASH_COLUMNS) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = STRIPE_FILL

        # Score cell styling
        score_cell = ws.cell(row=row_idx, column=1)
        score_cell.alignment = Alignment(horizontal="center", vertical="top")
        if score >= 65:
            score_cell.fill = SCORE_GREEN
            score_cell.font = SCORE_GREEN_FONT
        elif score >= 45:
            score_cell.fill = SCORE_YELLOW
            score_cell.font = SCORE_YELLOW_FONT
        else:
            score_cell.fill = SCORE_RED
            score_cell.font = SCORE_RED_FONT

        # Title as hyperlink
        title_cell = ws.cell(row=row_idx, column=2)
        url = opp.get("url", "")
        if url:
            title_cell.hyperlink = url
            title_cell.font = LINK_FONT
        title_cell.alignment = WRAP_ALIGN

        # SAM link column as clickable "View"
        link_cell = ws.cell(row=row_idx, column=len(DASH_COLUMNS))
        if url:
            link_cell.value = "View on SAM.gov"
            link_cell.hyperlink = url
            link_cell.font = LINK_FONT
            link_cell.alignment = Alignment(horizontal="center", vertical="top")

        # Days-left coloring (column 10 after adding First Seen)
        if days_left is not None:
            dl_cell = ws.cell(row=row_idx, column=10)
            dl_cell.alignment = Alignment(horizontal="center", vertical="top")
            if days_left < 0:
                dl_cell.font = Font(name="Calibri", size=10, color="9C0006")
            elif days_left <= 7:
                dl_cell.font = Font(name="Calibri", size=10, bold=True, color="9C6500")

    # Auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(DASH_COLUMNS))}{len(opportunities) + 1}"


# ------------------------------------------------------------------ details
DETAIL_COLUMNS = [
    ("Score",              6),
    ("Title",              48),
    ("Agency",             26),
    ("Office",             26),
    ("Type",               20),
    ("Set-Aside",          20),
    ("NAICS",              8),
    ("Posted",             11),
    ("First Seen",         11),
    ("Deadline",           11),
    ("Days Left",          9),
    ("Est. Value",         14),
    ("Location",           18),
    ("Contact",            30),
    ("Description",        60),
    ("Flags",              36),
    ("Score Breakdown",    40),
    ("GO/NO-GO",           8),
    ("AI Summary",         50),
    ("AI Requirements",    40),
    ("AI Feasibility",     40),
    ("AI Red Flags",       40),
    ("AI Recommendation",  40),
    ("SAM.gov Link",       14),
]


def _build_details(wb, opportunities):
    ws = wb.create_sheet("Details")
    ws.sheet_properties.tabColor = "2E75B6"

    # Header row
    for col_idx, (header, width) in enumerate(DETAIL_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

    now = datetime.now()
    for row_idx, opp in enumerate(opportunities, 2):
        score = opp.get("score", 0)
        deadline_str = opp.get("response_deadline") or ""
        days_left = _days_until(deadline_str, now)
        flags_text = "\n".join(opp.get("flags") or [])
        breakdown_text = _format_breakdown(opp.get("score_breakdown") or {})

        values = [
            score,
            opp.get("title", ""),
            opp.get("agency", ""),
            opp.get("office", ""),
            opp.get("type", ""),
            opp.get("set_aside", ""),
            opp.get("naics_code", ""),
            _format_date_short(opp.get("posted_date", "")),
            _format_date_short(opp.get("first_seen", "")),
            _format_date_short(deadline_str),
            days_left if days_left is not None else "",
            opp.get("estimated_value", ""),
            opp.get("place_of_performance", ""),
            opp.get("contact", ""),
            opp.get("description", ""),
            flags_text,
            breakdown_text,
            opp.get("ai_go_no_go", ""),
            opp.get("ai_summary", ""),
            opp.get("ai_requirements", ""),
            opp.get("ai_feasibility", ""),
            opp.get("ai_red_flags", ""),
            opp.get("ai_recommendation", ""),
            opp.get("url", ""),
        ]

        for col_idx, val in enumerate(values, 1):
            # Strip non-printable characters that Excel rejects
            if isinstance(val, str):
                val = "".join(c for c in val if c >= " " or c in "\t\n\r")
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = BODY_FONT
            cell.alignment = WRAP_ALIGN
            cell.border = THIN_BORDER

        # Alternate row shading
        if row_idx % 2 == 0:
            for col_idx in range(1, len(DETAIL_COLUMNS) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = STRIPE_FILL

        # Score styling
        score_cell = ws.cell(row=row_idx, column=1)
        score_cell.alignment = Alignment(horizontal="center", vertical="top")
        if score >= 65:
            score_cell.fill = SCORE_GREEN
            score_cell.font = SCORE_GREEN_FONT
        elif score >= 45:
            score_cell.fill = SCORE_YELLOW
            score_cell.font = SCORE_YELLOW_FONT
        else:
            score_cell.fill = SCORE_RED
            score_cell.font = SCORE_RED_FONT

        # Title as hyperlink
        title_cell = ws.cell(row=row_idx, column=2)
        url = opp.get("url", "")
        if url:
            title_cell.hyperlink = url
            title_cell.font = LINK_FONT
        title_cell.alignment = WRAP_ALIGN

        # GO/NO-GO cell styling
        gng_col = next((i + 1 for i, (h, _) in enumerate(DETAIL_COLUMNS) if h == "GO/NO-GO"), None)
        if gng_col:
            gng_cell = ws.cell(row=row_idx, column=gng_col)
            gng_val = opp.get("ai_go_no_go", "")
            gng_cell.alignment = Alignment(horizontal="center", vertical="top")
            if gng_val == "GO":
                gng_cell.fill = SCORE_GREEN
                gng_cell.font = SCORE_GREEN_FONT
            elif gng_val == "NO-GO":
                gng_cell.fill = SCORE_RED
                gng_cell.font = SCORE_RED_FONT

        # Link column
        link_cell = ws.cell(row=row_idx, column=len(DETAIL_COLUMNS))
        if url:
            link_cell.value = "View on SAM.gov"
            link_cell.hyperlink = url
            link_cell.font = LINK_FONT
            link_cell.alignment = Alignment(horizontal="center", vertical="top")

    ws.auto_filter.ref = f"A1:{get_column_letter(len(DETAIL_COLUMNS))}{len(opportunities) + 1}"


# ------------------------------------------------------------------ helpers
def _days_until(deadline_str, now=None):
    """Parse deadline and return days remaining, or None."""
    if not deadline_str:
        return None
    now = now or datetime.now()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%S%z"):
        try:
            dl = datetime.strptime(deadline_str[:19], fmt[:19] if "%z" in fmt else fmt)
            return (dl - now).days
        except ValueError:
            continue
    return None


def _format_date_short(date_str):
    """Turn '2026-02-10' or '02/10/2026' into 'Feb 10'."""
    if not date_str:
        return ""
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%Y-%m-%dT%H:%M:%S"):
        try:
            dt = datetime.strptime(date_str[:19], fmt)
            return dt.strftime("%b %d")
        except ValueError:
            continue
    return date_str


def _format_breakdown(breakdown):
    """Turn score_breakdown dict into readable multi-line text."""
    if not breakdown:
        return ""

    # New 4-category format
    categories = [
        ("entry_barrier",     "Entry Barrier",     40),
        ("scope_simplicity",  "Scope Simplicity",  25),
        ("competition_level", "Competition",        20),
        ("skill_match",       "Skill Match",        15),
    ]
    if any(k in breakdown for k, _, _ in categories):
        lines = []
        for key, label, max_pts in categories:
            cat = breakdown.get(key)
            if not isinstance(cat, dict) or "points" not in cat:
                continue
            pts = cat["points"]
            lines.append(f"{label}: {pts}/{max_pts}")
            for d in cat.get("details", []):
                dp = d.get("points", 0)
                if dp == 0:
                    continue
                sign = "+" if dp > 0 else ""
                matches = f" ({', '.join(d['matches'][:3])})" if d.get("matches") else ""
                lines.append(f"  {sign}{dp}  {d.get('label', '')}{matches}")
        if breakdown.get("rationale"):
            lines.append("")
            lines.append(breakdown["rationale"])
        return "\n".join(lines)

    # Legacy flat format fallback
    lines = []
    for category, info in breakdown.items():
        if category in ("rationale", "disqualifiers", "warnings"):
            continue
        if not isinstance(info, dict):
            continue
        pts = info.get("points", 0)
        if pts == 0:
            continue
        label = category.replace("_", " ").title()
        detail = ""
        if info.get("matches"):
            detail = f" ({', '.join(info['matches'][:3])})"
        elif info.get("value"):
            detail = f" ({info['value']})"
        sign = "+" if pts > 0 else ""
        lines.append(f"{label}: {sign}{pts}{detail}")
    return "\n".join(lines)
